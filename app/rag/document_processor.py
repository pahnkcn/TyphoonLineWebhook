"""Document loading and chunking utilities for API-based RAG."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Union

_TOKEN_PATTERN = re.compile(r"[A-Za-z0-9_]+|[\u0E00-\u0E7F]")
_SENTENCE_SPLIT_PATTERN = re.compile(r"(?<=[.!?。！？])\s+")
_MARKDOWN_HEADING_PATTERN = re.compile(r"(?m)^(#{1,6})\s+(.+)$")
_MARKDOWN_IMAGE_PATTERN = re.compile(r"!\[[^\]]*\]\([^)]+\)")
_MARKDOWN_LINK_PATTERN = re.compile(r"\[([^\]]+)\]\([^)]+\)")
_NOISE_LINE_PATTERNS = (
    re.compile(r"^\s*source\s*:\s*", flags=re.IGNORECASE),
    re.compile(r"^\s*table of contents\s*$", flags=re.IGNORECASE),
    re.compile(r"^\s*contents\s*$", flags=re.IGNORECASE),
    re.compile(r"^\s*copyright\b", flags=re.IGNORECASE),
    re.compile(r"^\s*all rights reserved\b", flags=re.IGNORECASE),
    re.compile(r"^\s*library of congress\b", flags=re.IGNORECASE),
    re.compile(r"^\s*a cip record\b", flags=re.IGNORECASE),
    re.compile(r"^\s*isbn\b", flags=re.IGNORECASE),
    re.compile(r"^\s*printed in\b", flags=re.IGNORECASE),
    re.compile(r"^\s*cover design\s*:", flags=re.IGNORECASE),
    re.compile(r"^\s*(?:https?://|www\.)", flags=re.IGNORECASE),
    re.compile(r"^\s*[ivxlcdm]+\s*$", flags=re.IGNORECASE),
    re.compile(r"^\s*\d+\s*$"),
    re.compile(r"^\s*[^A-Za-z\u0E00-\u0E7F]{3,}\s*$"),
)
_NOISE_SECTION_TITLES = {
    "about the authors",
    "acknowledgements",
    "acknowledgments",
    "bibliography",
    "contents",
    "contributors",
    "copyright",
    "foreword",
    "index",
    "library of congress cataloging in publication data",
    "preface",
    "table of contents",
}


@dataclass
class DocumentChunk:
    """A single chunk of knowledge content with searchable metadata."""

    content: str
    metadata: Dict[str, object]


@dataclass
class DocumentSegment:
    """A semantic section of a document before chunking."""

    content: str
    metadata: Dict[str, object]


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in str(text or "").splitlines()]
    compact = "\n".join(line for line in lines if line)
    return compact.strip()


def _normalize_title(text: str) -> str:
    lowered = str(text or "").strip().lower()
    lowered = re.sub(r"[^a-z0-9\s]", " ", lowered)
    return re.sub(r"\s+", " ", lowered).strip()


def _strip_markdown_artifacts(text: str) -> str:
    stripped = str(text or "").replace("\u00a0", " ")
    stripped = re.sub(r"&nbsp;?", " ", stripped, flags=re.IGNORECASE)
    stripped = _MARKDOWN_IMAGE_PATTERN.sub(" ", stripped)
    stripped = _MARKDOWN_LINK_PATTERN.sub(r"\1", stripped)
    return stripped


def _is_noise_title(title: str) -> bool:
    return _normalize_title(title) in _NOISE_SECTION_TITLES


def _is_noise_line(line: str) -> bool:
    candidate = str(line or "").strip()
    if not candidate:
        return True
    if any(pattern.search(candidate) for pattern in _NOISE_LINE_PATTERNS):
        return True
    if re.search(r"\.{3,}\s*(?:[A-Za-z0-9]+|[ivxlcdm]+)\s*$", candidate, flags=re.IGNORECASE):
        return True
    return False


def _sanitize_text(text: str) -> str:
    stripped = _strip_markdown_artifacts(text)
    cleaned_lines: List[str] = []
    for raw_line in stripped.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if _is_noise_line(line):
            continue
        cleaned_lines.append(line)
    return _normalize_text("\n".join(cleaned_lines))


def _is_low_signal_chunk(content: str) -> bool:
    cleaned = _normalize_text(content)
    if not cleaned:
        return True
    if _is_noise_line(cleaned):
        return True
    token_count = _token_count(cleaned)
    text_char_count = len(re.findall(r"[A-Za-z\u0E00-\u0E7F]", cleaned))
    return token_count < 6 and text_char_count < 24


def _sanitize_segments(segments: Sequence[DocumentSegment]) -> List[DocumentSegment]:
    cleaned_segments: List[DocumentSegment] = []
    for segment in segments:
        metadata = dict(segment.metadata or {})
        section_title = str(metadata.get("section_title") or "").strip()
        if section_title and _is_noise_title(section_title):
            continue
        content = _sanitize_text(segment.content)
        if not content or _is_low_signal_chunk(content):
            continue
        if section_title:
            metadata["section_title"] = _sanitize_text(section_title) or section_title
        cleaned_segments.append(DocumentSegment(content=content, metadata=metadata))
    return cleaned_segments


def _detect_language(text: str) -> str:
    thai_count = len(re.findall(r"[\u0E00-\u0E7F]", text))
    latin_count = len(re.findall(r"[A-Za-z]", text))
    if thai_count > 0 and latin_count > 0:
        return "mixed"
    if thai_count > 0:
        return "th"
    if latin_count > 0:
        return "en"
    return "unknown"


def _token_count(text: str) -> int:
    if not text:
        return 0
    return len(_TOKEN_PATTERN.findall(text))


def _join_segments(segments: Sequence[DocumentSegment]) -> str:
    return _normalize_text("\n\n".join(segment.content for segment in segments if segment.content.strip()))


def _split_markdown_sections(text: str) -> List[DocumentSegment]:
    cleaned = _normalize_text(text)
    if not cleaned:
        return []

    matches = list(_MARKDOWN_HEADING_PATTERN.finditer(cleaned))
    if not matches:
        return [
            DocumentSegment(
                content=cleaned,
                metadata={"section_title": "", "segment_type": "text"},
            )
        ]

    segments: List[DocumentSegment] = []
    preamble = cleaned[: matches[0].start()].strip()
    if preamble:
        segments.append(
            DocumentSegment(
                content=preamble,
                metadata={"section_title": "", "segment_type": "markdown"},
            )
        )

    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(cleaned)
        block = cleaned[start:end].strip()
        if not block:
            continue
        title = match.group(2).strip()
        segments.append(
            DocumentSegment(
                content=block,
                metadata={"section_title": title, "segment_type": "markdown"},
            )
        )
    return segments


def load_pdf_segments(file_path: Union[Path, str]) -> List[DocumentSegment]:
    file_path = Path(file_path)
    from pypdf import PdfReader  # lazy import to keep startup lightweight

    reader = PdfReader(str(file_path))
    segments: List[DocumentSegment] = []
    for page_index, page in enumerate(reader.pages, start=1):
        extracted = _normalize_text(page.extract_text() or "")
        if not extracted:
            continue
        segments.append(
            DocumentSegment(
                content=extracted,
                metadata={
                    "page_start": page_index,
                    "page_end": page_index,
                    "section_title": f"page_{page_index}",
                    "segment_type": "pdf_page",
                },
            )
        )
    return segments


def load_docx_segments(file_path: Union[Path, str]) -> List[DocumentSegment]:
    file_path = Path(file_path)
    from docx import Document  # lazy import to keep startup lightweight

    doc = Document(str(file_path))
    sections: List[DocumentSegment] = []
    current_title = ""
    current_lines: List[str] = []

    def _flush_current_section() -> None:
        nonlocal current_lines
        content = _normalize_text("\n\n".join(current_lines))
        if content:
            sections.append(
                DocumentSegment(
                    content=content,
                    metadata={
                        "section_title": current_title,
                        "segment_type": "docx_section",
                    },
                )
            )
        current_lines = []

    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue

        style_name = ""
        try:
            style_name = str((paragraph.style.name if paragraph.style else "") or "").lower()
        except Exception:
            style_name = ""

        is_heading = style_name.startswith("heading")
        if is_heading:
            _flush_current_section()
            current_title = text
            continue

        current_lines.append(text)

    _flush_current_section()
    if sections:
        return sections

    fallback = _normalize_text("\n\n".join((p.text or "").strip() for p in doc.paragraphs if (p.text or "").strip()))
    if not fallback:
        return []
    return [
        DocumentSegment(
            content=fallback,
            metadata={"section_title": "", "segment_type": "docx"},
        )
    ]


def load_text_segments(file_path: Union[Path, str]) -> List[DocumentSegment]:
    file_path = Path(file_path)
    raw_text = file_path.read_text(encoding="utf-8", errors="ignore")
    if file_path.suffix.lower() == ".md":
        return _split_markdown_sections(raw_text)

    cleaned = _normalize_text(raw_text)
    if not cleaned:
        return []
    return [
        DocumentSegment(
            content=cleaned,
            metadata={"section_title": "", "segment_type": "text"},
        )
    ]


def load_document_segments(file_path: Union[Path, str]) -> List[DocumentSegment]:
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        segments = load_pdf_segments(file_path)
    elif suffix == ".docx":
        segments = load_docx_segments(file_path)
    elif suffix in {".txt", ".md"}:
        segments = load_text_segments(file_path)
    else:
        raise ValueError(f"Unsupported document type: {file_path.name}")
    return _sanitize_segments(segments)


def load_pdf(file_path: Union[Path, str]) -> str:
    return _join_segments(load_document_segments(file_path))


def load_docx(file_path: Union[Path, str]) -> str:
    return _join_segments(load_document_segments(file_path))


def load_text(file_path: Union[Path, str]) -> str:
    return _join_segments(load_document_segments(file_path))


def load_document(file_path: Union[Path, str]) -> str:
    return _join_segments(load_document_segments(file_path))


def _recursive_split(text: str, max_len: int, separators: Iterable[str]) -> List[str]:
    if len(text) <= max_len:
        return [text]

    for sep in separators:
        if sep not in text:
            continue

        pieces = text.split(sep)
        chunks: List[str] = []
        current = ""

        for piece in pieces:
            part = piece.strip()
            if not part:
                continue

            candidate = part if not current else f"{current}{sep}{part}"
            if len(candidate) <= max_len:
                current = candidate
            else:
                if current:
                    chunks.append(current.strip())
                if len(part) > max_len:
                    chunks.extend(_recursive_split(part, max_len, separators))
                    current = ""
                else:
                    current = part

        if current:
            chunks.append(current.strip())

        if chunks:
            return chunks

    return [text[i : i + max_len].strip() for i in range(0, len(text), max_len)]


def _split_by_headings(text: str, chunk_size: int) -> List[str]:
    sections = re.split(r"(?m)^(?=## )", text)
    sections = [section.strip() for section in sections if section.strip()]
    if len(sections) <= 1:
        return []

    chunks: List[str] = []
    current = ""
    for section in sections:
        candidate = f"{current}\n\n{section}".strip() if current else section
        if len(candidate) <= chunk_size:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(section) > chunk_size:
            chunks.extend(_recursive_split(section, chunk_size, ["\n\n", "\n", ". ", " "]))
            current = ""
        else:
            current = section

    if current:
        chunks.append(current)
    return [chunk for chunk in chunks if chunk]


def _chunk_text_char_mode(text: str, chunk_size: int, overlap: int) -> List[str]:
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    heading_chunks = _split_by_headings(text, chunk_size)
    if heading_chunks:
        base_chunks = heading_chunks
    else:
        base_chunks = _recursive_split(text, chunk_size, ["\n\n", "\n", ". ", " "])

    if overlap <= 0:
        return [chunk for chunk in base_chunks if chunk]

    merged: List[str] = []
    for index, chunk in enumerate(base_chunks):
        chunk = chunk.strip()
        if not chunk:
            continue
        if index == 0:
            merged.append(chunk)
            continue
        prev_source = base_chunks[index - 1].strip()
        prev_tail = prev_source[-overlap:] if len(prev_source) > overlap else prev_source
        merged.append(f"{prev_tail} {chunk}".strip())
    return merged


def _resolve_token_limit(chunk_size: int) -> int:
    if chunk_size <= 0:
        return 200
    return max(120, min(900, chunk_size // 4))


def _resolve_overlap_tokens(overlap: int) -> int:
    if overlap <= 0:
        return 0
    return max(20, min(220, overlap // 4 if overlap > 80 else overlap))


def _tail_for_overlap(text: str, overlap_tokens: int) -> str:
    if overlap_tokens <= 0:
        return ""
    if " " in text:
        words = text.split()
        if len(words) <= overlap_tokens:
            return " ".join(words)
        return " ".join(words[-overlap_tokens:])
    fallback_chars = max(60, overlap_tokens * 3)
    return text[-fallback_chars:]


def _split_into_units(text: str) -> List[str]:
    if not text:
        return []

    markdown_sections = re.split(r"(?m)(?=^#{1,6}\s+)", text)
    markdown_units = [unit.strip() for unit in markdown_sections if unit and unit.strip()]
    if len(markdown_units) > 1:
        return markdown_units

    paragraphs = [paragraph.strip() for paragraph in re.split(r"\n{2,}", text) if paragraph.strip()]
    if len(paragraphs) > 1:
        return paragraphs

    sentences = [sentence.strip() for sentence in _SENTENCE_SPLIT_PATTERN.split(text) if sentence.strip()]
    if len(sentences) > 1:
        return sentences

    return [text.strip()]


def _split_long_unit(unit: str, token_limit: int) -> List[str]:
    if _token_count(unit) <= token_limit:
        return [unit]

    sentences = [sentence.strip() for sentence in _SENTENCE_SPLIT_PATTERN.split(unit) if sentence.strip()]
    if len(sentences) > 1:
        chunks: List[str] = []
        current: List[str] = []
        current_tokens = 0
        for sentence in sentences:
            sentence_tokens = max(1, _token_count(sentence))
            if current and current_tokens + sentence_tokens > token_limit:
                chunks.append(" ".join(current).strip())
                current = [sentence]
                current_tokens = sentence_tokens
            else:
                current.append(sentence)
                current_tokens += sentence_tokens
        if current:
            chunks.append(" ".join(current).strip())
        if chunks:
            return [chunk for chunk in chunks if chunk]

    char_window = max(320, token_limit * 4)
    return [unit[index : index + char_window].strip() for index in range(0, len(unit), char_window) if unit[index : index + char_window].strip()]


def _chunk_text_token_mode(text: str, chunk_size: int, overlap: int) -> List[str]:
    token_limit = _resolve_token_limit(chunk_size)
    overlap_tokens = _resolve_overlap_tokens(overlap)

    if _token_count(text) <= token_limit:
        return [text]

    units = _split_into_units(text)
    prepared_units: List[str] = []
    for unit in units:
        prepared_units.extend(_split_long_unit(unit, token_limit))

    chunks: List[str] = []
    current_units: List[str] = []
    current_tokens = 0

    for unit in prepared_units:
        unit_tokens = max(1, _token_count(unit))
        if current_units and current_tokens + unit_tokens > token_limit:
            chunk_text_value = "\n\n".join(current_units).strip()
            if chunk_text_value:
                chunks.append(chunk_text_value)

            if overlap_tokens > 0 and chunk_text_value:
                overlap_text = _tail_for_overlap(chunk_text_value, overlap_tokens)
                current_units = [overlap_text] if overlap_text else []
                current_tokens = _token_count(overlap_text)
            else:
                current_units = []
                current_tokens = 0

        current_units.append(unit)
        current_tokens += unit_tokens

    if current_units:
        chunk_text_value = "\n\n".join(current_units).strip()
        if chunk_text_value:
            chunks.append(chunk_text_value)

    return chunks


def _same_page_span(left_meta: Dict[str, object], right_meta: Dict[str, object]) -> bool:
    return (
        left_meta.get("page_start") == right_meta.get("page_start")
        and left_meta.get("page_end") == right_meta.get("page_end")
    )


def _merge_section_titles(left_title: str, right_title: str) -> str:
    left_clean = str(left_title or "").strip()
    right_clean = str(right_title or "").strip()
    if not left_clean:
        return right_clean
    if not right_clean or right_clean == left_clean:
        return left_clean
    return f"{left_clean} | {right_clean}"


def _merge_chunk_metadata(left_meta: Dict[str, object], right_meta: Dict[str, object]) -> Dict[str, object]:
    merged = dict(left_meta or {})
    merged["section_title"] = _merge_section_titles(
        str(left_meta.get("section_title") or ""),
        str(right_meta.get("section_title") or ""),
    )
    merged["page_start"] = left_meta.get("page_start", right_meta.get("page_start"))
    merged["page_end"] = right_meta.get("page_end", left_meta.get("page_end"))
    merged["segment_index_end"] = right_meta.get("segment_index")
    return merged


def _merge_short_document_chunks(
    chunks: List[DocumentChunk],
    min_tokens: int,
    max_tokens: int,
) -> List[DocumentChunk]:
    if not chunks:
        return []

    merged: List[DocumentChunk] = []
    for chunk in chunks:
        content = str(chunk.content or "").strip()
        if not content:
            continue

        current = DocumentChunk(content=content, metadata=dict(chunk.metadata or {}))
        current_tokens = _token_count(content)

        if merged and current_tokens < min_tokens:
            previous = merged[-1]
            previous_tokens = _token_count(previous.content)
            if (
                previous_tokens + current_tokens <= max_tokens
                and _same_page_span(previous.metadata, current.metadata)
            ):
                merged[-1] = DocumentChunk(
                    content=f"{previous.content}\n\n{current.content}".strip(),
                    metadata=_merge_chunk_metadata(previous.metadata, current.metadata),
                )
                continue

        merged.append(current)

    if len(merged) >= 2:
        last = merged[-1]
        previous = merged[-2]
        last_tokens = _token_count(last.content)
        previous_tokens = _token_count(previous.content)
        if (
            last_tokens < min_tokens
            and previous_tokens + last_tokens <= max_tokens
            and _same_page_span(previous.metadata, last.metadata)
        ):
            merged[-2] = DocumentChunk(
                content=f"{previous.content}\n\n{last.content}".strip(),
                metadata=_merge_chunk_metadata(previous.metadata, last.metadata),
            )
            merged.pop()

    return merged


def chunk_text(text: str, chunk_size: int = 1100, overlap: int = 120) -> List[str]:
    if not text or not text.strip():
        return []

    cleaned = text.strip()
    if chunk_size <= 0:
        chunk_size = 1100
    if overlap < 0:
        overlap = 0

    # Keep character-based behavior for very small chunk sizes to preserve legacy behavior.
    if chunk_size <= 256:
        return _chunk_text_char_mode(cleaned, chunk_size=chunk_size, overlap=overlap)

    return _chunk_text_token_mode(cleaned, chunk_size=chunk_size, overlap=overlap)


def chunk_document(
    file_path: str,
    chunk_size: int = 1100,
    overlap: int = 120,
) -> List[DocumentChunk]:
    path = Path(file_path)
    segments = load_document_segments(path)

    provisional_chunks: List[DocumentChunk] = []

    for segment_index, segment in enumerate(segments):
        base_metadata = dict(segment.metadata or {})
        section_chunks = chunk_text(segment.content, chunk_size=chunk_size, overlap=overlap)

        for chunk_in_segment, chunk_content in enumerate(section_chunks):
            content = chunk_content.strip()
            if not content:
                continue
            metadata = {
                "source": path.name,
                "source_path": str(path.resolve()),
                "chunk_in_segment": chunk_in_segment,
                "segment_index": segment_index,
                "section_title": str(base_metadata.get("section_title") or ""),
                "page_start": base_metadata.get("page_start"),
                "page_end": base_metadata.get("page_end"),
            }
            provisional_chunks.append(DocumentChunk(content=content, metadata=metadata))

    if chunk_size > 256 and provisional_chunks:
        token_limit = _resolve_token_limit(chunk_size)
        min_tokens = max(24, min(80, token_limit // 8))
        max_tokens = max(min_tokens * 3, int(token_limit * 0.90))
        provisional_chunks = _merge_short_document_chunks(
            provisional_chunks,
            min_tokens=min_tokens,
            max_tokens=max_tokens,
        )

    results: List[DocumentChunk] = []
    for chunk in provisional_chunks:
        content = str(chunk.content or "").strip()
        if not content or _is_low_signal_chunk(content):
            continue
        token_count = _token_count(content)
        metadata = dict(chunk.metadata or {})
        metadata.update(
            {
                "chunk_index": len(results),
                "language": _detect_language(content),
                "char_count": len(content),
                "token_count": token_count,
            }
        )
        results.append(DocumentChunk(content=content, metadata=metadata))

    return results
